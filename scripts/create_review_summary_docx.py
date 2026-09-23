import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

def create_review_docx():
    doc = docx.Document()
    
    # 1. Page Setup - Margins 1 inch
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Styling Helpers
    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def add_callout(title, text, border_color="1B365D", fill_color="F4F6F9"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, fill_color)
        set_cell_margins(cell, top=110, bottom=110, left=150, right=140)
        
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
        tcPr.append(tcBorders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        rt = p.add_run(f"★ {title}\n")
        rt.bold = True
        rt.font.name = "Arial"
        rt.font.size = Pt(10.5)
        rt.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        
        rb = p.add_run(text)
        rb.font.name = "Arial"
        rb.font.size = Pt(9.5)
        rb.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(13.5)
        r.bold = True
        r.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(11.0)
        r.bold = True
        r.font.color.rgb = RGBColor(0x00, 0x5A, 0x9E)
        return p

    def add_body(text, bold_prefix=None, space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.bold = True
            rb.font.name = "Arial"
            rb.font.size = Pt(9.5)
            rb.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.bold = True
            rb.font.name = "Arial"
            rb.font.size = Pt(9.5)
            rb.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        return p

    def add_code_block(code_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, "282C34")
        set_cell_margins(cell, top=70, bottom=70, left=110, right=110)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.1
        r = p.add_run(code_text)
        r.font.name = "Courier New"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0xAB, 0xB2, 0xBF)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def create_table(col_widths, headers, data, header_bg="1B365D", alt_bg="F9FAFB"):
        tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        
        # Header row
        hdr_row = tbl.rows[0]
        for i, h_text in enumerate(headers):
            cell = hdr_row.cells[i]
            cell.width = col_widths[i]
            set_cell_background(cell, header_bg)
            set_cell_margins(cell, top=70, bottom=70, left=60, right=60)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(h_text)
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
        for r_idx, row_data in enumerate(data):
            row = tbl.rows[r_idx + 1]
            bg = alt_bg if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, val in enumerate(row_data):
                cell = row.cells[c_idx]
                cell.width = col_widths[c_idx]
                set_cell_background(cell, bg)
                set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                
                if c_idx in [0] and len(headers) >= 4:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    
                run = p.add_run(str(val))
                run.font.name = "Arial"
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
                if c_idx == 0:
                    run.bold = True
        doc.add_paragraph().paragraph_format.space_after = Pt(5)

    # ==========================================
    # HEADER & TITLE
    # ==========================================
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(10)
    title_p.paragraph_format.space_after = Pt(3)
    run_main = title_p.add_run("BẢN ĐÁNH GIÁ CHUYÊN MÔN & TỔNG KẾT KỸ THUẬT NCKH\n")
    run_main.bold = True
    run_main.font.name = "Arial"
    run_main.font.size = Pt(16)
    run_main.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    
    run_sub = title_p.add_run("HOÀN THIỆN PHƯƠNG ÁN XỬ LÝ KHUYẾT THIẾU MODALITY & BẢN THỂ HỌC GENRE NHẠC VIỆT NAM\n(Multimodal Music Genre Classification Under Missing Modalities & Multi-Label)")
    run_sub.bold = True
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(0x00, 0x5A, 0x9E)
    
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_before = Pt(3)
    meta_p.paragraph_format.space_after = Pt(10)
    r_meta = meta_p.add_run("Người đánh giá: AI Specialist & Senior Data Engineer | Dự án: DatasetMusic\nMục tiêu: Đề tài Nghiên cứu Khoa học (NCKH) & Bảo vệ Hội đồng Chuyên môn | GPU: RTX 3050 Ti (4GB)")
    r_meta.italic = True
    r_meta.font.name = "Arial"
    r_meta.font.size = Pt(9.0)
    r_meta.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # ==========================================
    # PHẦN 1: ĐÁNH GIÁ TỔNG QUAN
    # ==========================================
    add_heading_1("PHẦN 1: ĐÁNH GIÁ TỔNG QUAN BẢN CẬP NHẬT MỚI NHẤT (ĐIỂM SỐ: 9.0 / 10)")
    
    add_callout("KẾT LUẬN ĐÁNH GIÁ CHUYÊN MÔN",
        "Bản cập nhật mới nhất đã đạt được độ mạch lạc cao, phản ánh đúng bản chất bài toán và thoát khỏi các lối mòn học thuật đơn giản. "
        "Việc bổ sung ý tưởng 'Mỗi đặc trưng gắn Confidence' là một điểm sáng xuất sắc. "
        "Với 3 hiệu chỉnh nhỏ: (1) Đổi ngưỡng cố định 95% thành Adaptive Threshold, (2) Nêu rõ lý do kỹ thuật phản biện, và (3) Gọi tên chuẩn xác các phương pháp SOTA (AudioMAE, Gated Fusion, ASL), "
        "câu trả lời hoàn toàn đạt tiêu chuẩn xuất bản và bảo vệ xuất sắc trước hội đồng NCKH.")

    add_heading_2("Điểm mạnh nổi bật:")
    add_bullet(" Mạch lạc, ngắn gọn, đi thẳng vào trọng tâm cốt lõi của bài toán dữ liệu thực tế.", "Tư duy giải pháp: ");
    add_bullet(" Rất chuẩn xác, giúp mô hình biết linh hoạt điều chỉnh tỷ trọng đóng góp của từng phương thức.", "Thêm cơ chế Confidence per feature: ");
    add_bullet(" Không chạy theo lối mòn gọi API bừa bãi hay ghép nhãn ca sĩ gây rò rỉ dữ liệu.", "Tư duy phản biện xuất sắc với Spotify / Ghép chéo: ");

    add_heading_2("Ba điểm mấu chốt cần chỉnh để đạt độ chặt chẽ tuyệt đối:")
    add_bullet(" Ngưỡng cố định sẽ xóa sổ các thể loại hiếm (Nhạc Đỏ, Cổ truyền). Bắt buộc đổi sang Ngưỡng thích ứng theo từng lớp (Adaptive Threshold 0.68 – 0.88).", "1. Sửa ngưỡng 95% cố định: ");
    add_bullet(" Làm rõ tại sao cấm Whisper trên nhạc cổ truyền, tại sao Spotify lệch domain, tại sao ghép chéo gây data leakage.", "2. Bổ sung lý do cho từng quyết định: ");
    add_bullet(" Gọi tên chính thức: Self-Supervised Pretraining (AudioMAE/CLAP), Gated Cross-Attention, Asymmetric Loss (ASL).", "3. Chuẩn hóa thuật ngữ khoa học: ");

    # ==========================================
    # PHẦN 2: CHI TIẾT 5 Ý VỀ THIẾU MODALITY
    # ==========================================
    add_heading_1("PHẦN 2: PHÂN TÍCH CHUYÊN SÂU 5 Ý VỀ DỮ LIỆU THIẾU MODALITY")

    add_heading_2("Ý 1: Tách lời + Gắn cờ + Giảm Confidence")
    add_body("Ý bạn nêu: 'Dùng tool tách lời (nhưng vẫn phải check lại), gắn cờ và giảm độ tin tưởng khi kết hợp với các dữ liệu có sẵn.' -> Đúng và đủ ý.", "Đánh giá: ");
    add_body("Bổ sung 2 điều kiện tiên quyết bắt buộc phải có để câu trả lời chặt chẽ trước hội đồng:");
    add_bullet("Tuyệt đối không chạy Whisper trực tiếp trên tệp nhạc hỗn hợp có beat, trống bass và hiệu ứng reverb/autotune. Tiếng nhạc nền sẽ khiến Whisper bị ảo giác loop vô tận ('Cảm ơn các bạn...'). Phải dùng Demucs v4 hoặc MDX-Net bóc tách vocal trước.", "Điều kiện 1 — Bắt buộc tách vocal trước ASR: ");
    add_bullet("Ca từ cổ ngữ, tiết tấu luyến láy ngân nga và tiếng đàn đáy, đàn nhị, phách cổ truyền lấn át hoàn toàn vocal. Whisper hiện tại sinh rác 100% trên dòng nhạc này. Phải loại trừ hoàn toàn nhóm này khỏi ASR.", "Điều kiện 2 — Tuyệt đối không dùng ASR cho Nhạc Cổ Truyền: ");

    add_heading_2("Ý 2: Gắn Mask = 0")
    add_body("Ý bạn nêu: 'Gắn mask_lyric, mask_audio = 0,...' -> Chuẩn mực SOTA trong Missing Modality Learning.", "Đánh giá: ");
    add_body("Bổ sung 1 điểm bản chất: Mask = 0 bắt buộc phải đi kèm với Gated Fusion, không chỉ zero-out vector đơn thuần. Nếu chỉ zero-out vector mà không có cơ chế Masked Gating, mô hình vẫn coi vector 0 là một tín hiệu đặc trưng hợp lệ, dẫn đến học sai lệch.", "Lưu ý then chốt: ");

    add_heading_2("Ý 3: Phân biệt trường hợp 'Chỉ có Avatar (Cover Art)'")
    add_body("Ý bạn nêu: 'Chỉ có avatar không thôi thì có thể bỏ luôn data đấy.' -> Rất đúng thực tế.", "Đánh giá: ");
    
    headers_t3 = ["Trường hợp dữ liệu thực tế", "Hành động kỹ thuật", "Lý do chuyên môn"]
    col_w_t3 = [Inches(2.0), Inches(1.8), Inches(2.7)]
    data_t3 = [
        ["Chỉ có cover, không audio, không lyric", "LOẠI BỎ KHỎI TRAIN", "Cover Art không mang đủ tín hiệu âm học để xác định thể loại đơn độc."],
        ["Có audio + cover, thiếu lyric", "GIỮ LẠI ĐỂ HUẤN LUYỆN", "Audio là phương thức mạnh nhất, chỉ cần gán mask_lyric = 0 là chạy tối ưu."]
    ]
    create_table(col_w_t3, headers_t3, data_t3, header_bg="005A9E")

    add_heading_2("Ý 4: Phản biện tại sao KHÔNG NÊN dùng API Spotify và Ghép chéo nguồn")
    add_body("Ý bạn nêu: 'Thấy 2 cách trong ngoặc không ổn cho lắm.' -> Trực giác Data Engineer rất chính xác! Cần nêu rõ 3 lý do học thuật:");
    add_bullet("Spotify không có kho dữ liệu đầy đủ cho nhạc truyền thống Việt Nam (Cải lương, Ca trù, Nhạc Đỏ). Siêu dữ liệu của Spotify sử dụng taxonomy phương Tây, không thể ánh xạ tương thích với âm nhạc Việt Nam.", "Lý do 1 — Lệch Domain & Taxonomy (Spotify): ");
    add_bullet("Một nghệ sĩ thể hiện rất nhiều thể loại khác nhau (Sơn Tùng M-TP hát Pop, R&B, Rap). Nếu ghép ảnh/nhãn theo ca sĩ, mô hình sẽ học nhận diện khuôn mặt ca sĩ thay vì nhận diện thể loại.", "Lý do 2 — Giả định sai lầm (Ghép chéo): ");
    add_bullet("Gây rò rỉ dữ liệu giữa tập Train và Test, tạo ra hiện tượng ảo giác dữ liệu (Data Hallucination).", "Lý do 3 — Data Leakage: ");

    add_heading_2("Ý 5: Cơ chế 'Confidence Feature' — Cụ thể hóa 4 Mục đích Sử dụng")
    add_body("Ý bạn nêu: 'Mỗi đặc trưng sẽ được gắn confidence để mô hình biết chọn.' -> Đây là ý tưởng rất xuất sắc, cần được cụ thể hóa thành 4 vai trò:");
    add_bullet("Trong Gated Cross-Attention, logit của từng modality được cộng thêm log(confidence). Phương thức nào có confidence thấp sẽ tự động bị giảm trọng số đóng góp.", "1. Điều chỉnh trọng số Fusion (Fusion Weighting): ");
    add_bullet("Mẫu có lyric sinh từ ASR (confidence = 0.6) sẽ có trọng số phạt loss nhỏ hơn mẫu có lyric chuẩn do con người kiểm duyệt (confidence = 1.0).", "2. Điều chỉnh trọng số hàm Loss (Sample Weighting): ");
    add_bullet("Khi tạo nhãn giả, chỉ những bài có confidence vượt ngưỡng mới được đưa vào tập huấn luyện bổ sung.", "3. Bộ lọc nhãn giả (Pseudo-Label Filtering): ");
    add_bullet("Những bản ghi có confidence nằm ở khoảng phân vân (0.4 – 0.6) được trích xuất riêng để chuyên gia con người thẩm định.", "4. Hỗ trợ kiểm duyệt con người (Human-in-the-loop Review): ");

    # Code block formula
    code_conf = """# Ví dụ: Modality Confidence tác động trực tiếp vào cơ chế Gating:
# h_m: vector biểu diễn, conf_m: điểm tin cậy thuộc [0.0, 1.0]
g_m = torch.matmul(w_m, h_m) + beta_m + torch.log(conf_m.clamp(min=1e-6))
# Nếu conf_m thấp -> g_m âm lớn -> alpha_m bị Softmax triệt tiêu về 0"""
    add_code_block(code_conf)

    # ==========================================
    # PHẦN 3: CHI TIẾT 3 Ý VỀ BÀI TOÁN GENRE
    # ==========================================
    add_heading_1("PHẦN 3: PHÂN TÍCH CHUYÊN SÂU 3 Ý VỀ BÀI TOÁN GENRE")

    add_heading_2("Ý 1: Học tự giám sát (SSL Pretraining) rồi Fine-tune")
    add_body("Ý bạn nêu: 'Dùng data không genre để train nhận diện nhạc Việt trước, sau đó fine-tune trên tập có nhãn.' -> Hoàn toàn chuẩn mực NCKH.", "Đánh giá: ");
    add_body("Cần gọi tên kỹ thuật chính thống: Đây là kỹ thuật Self-Supervised Learning (SSL) cho Audio. Các kiến trúc SOTA tiêu biểu gồm: AudioMAE (Masked Autoencoders for Audio), CLAP (Contrastive Language-Audio Pretraining), BYOL-A và SimCLR-Audio.", "Chuẩn hóa tên gọi: ");

    add_heading_2("Ý 2: Gán nhãn giả — Tại sao phải dùng Adaptive Threshold thay cho 95% cố định?")
    add_body("Ý bạn nêu: 'Chỉ lấy độ tin tưởng gần như rất cao 95%.' -> Đúng hướng thận trọng, nhưng ngưỡng cố định 95% cho tất cả các lớp sẽ gây ra lỗi nghiêm trọng: Các lớp hiếm (Nhạc Đỏ, Cổ truyền) có số lượng ít nên mạng dự đoán dè dặt, không bao giờ đạt được 95%, dẫn đến lớp hiếm bị xóa sổ hoàn toàn.", "Đánh giá & Phản biện: ");
    add_body("Giải pháp tối ưu: Ngưỡng thích ứng theo từng thể loại (Class-Adaptive Thresholds):", "Giải pháp SOTA: ");

    headers_t_ad = ["Nhóm thể loại", "Ngưỡng tin cậy chấp nhận (Threshold)", "Lý do kỹ thuật"]
    col_w_ad = [Inches(1.8), Inches(1.8), Inches(2.9)]
    data_t_ad = [
        ["Pop / Ballad", "P >= 0.88 (hoặc 0.90)", "Lớp chiếm đa số (45%), dễ gây nhầm lẫn nên phải siết chặt."],
        ["Hiphop / Rap / R&B", "P >= 0.85", "Lớp trung bình, tiết tấu và nhịp điệu đặc trưng."],
        ["Nhạc Đỏ / Cổ Truyền / Thiếu Nhi", "P >= 0.68 - 0.72", "Lớp thiểu số rất hiếm. Cần nới lỏng ngưỡng để thu gom thêm mẫu."]
    ]
    create_table(col_w_ad, headers_t_ad, data_t_ad, header_bg="005A9E")
    
    add_body("Kiểm soát chặt chẽ quy trình Pseudo-Labeling bằng bộ 4 điều kiện:");
    add_bullet("Chỉ lấy Top 50.000 bài có độ tự tin cao nhất trong 355k bài Zalo AI.", "1. Top-K Filter: ");
    add_bullet("Gán hệ số phạt loss lambda = 0.5 khi huấn luyện Student để tránh lây lan lỗi gán nhãn sai.", "2. Loss Weight: ");
    add_bullet("Áp dụng Label Smoothing = 0.1 để làm mềm phân phối xác suất mục tiêu.", "3. Label Smoothing: ");
    add_bullet("Lấy mẫu ngẫu nhiên 500 bài trong tập pseudo-label để người kiểm tra xác suất thực tế.", "4. Human Spot Check: ");

    add_heading_2("Ý 3: Chuyển đổi Multi-Class sang Multi-Label (Thay Softmax bằng Sigmoid)")
    add_body("Ý bạn nêu: 'Thay multi-class thành multi-label (thay softmax thành sigmoid).' -> Chuẩn mực toán học bắt buộc của bài toán âm nhạc.", "Đánh giá: ");
    add_bullet("Hàm Softmax ép tổng xác suất bằng 1, khiến các thể loại triệt tiêu lẫn nhau. Nhạc hiện đại thường là sự kết hợp: Pop + R&B, Rap + Pop, Rock + Nhạc Đỏ.", "Tại sao cấm Softmax: ");
    add_bullet("Tỷ lệ mẫu âm/dương lên tới 7:1. Bắt buộc dùng Asymmetric Loss (ASL) với cơ chế dịch biên xác suất để triệt tiêu gradient từ các nhãn âm dễ đoán.", "Hàm Loss bắt buộc (ASL): ");
    add_bullet("Không dùng ngưỡng mặc định 0.5. Quét tìm ngưỡng tối ưu riêng tau_c thuộc [0.2, 0.8] trên tập Validation nhằm tối đa hóa Macro F1.", "Adaptive Thresholding: ");

    # ==========================================
    # PHẦN 4: SO SÁNH 3 PHIÊN BẢN TIẾN HÓA
    # ==========================================
    add_heading_1("PHẦN 4: BẢNG SO SÁNH TIẾN HÓA NỘI DUNG QUA 3 PHIÊN BẢN")

    headers_comp = ["Tiêu chí kỹ thuật", "Bản 1 (Sơ khai)", "Bản 2 (Mở rộng)", "Bản 3 (Hoàn thiện NCKH)"]
    col_w_comp = [Inches(1.8), Inches(1.5), Inches(1.5), Inches(1.7)]
    data_comp = [
        ["Tách lời ASR", "ASR đơn thuần", "Gắn cờ + Giảm confidence", "Tách Vocal Demucs + Loại bỏ Cổ truyền"],
        ["Xử lý Modality Masking", "Zero-imputation", "Mask = 0", "Mask = 0 + Gated Cross-Attention"],
        ["Chỉ có Avatar", "Chưa phân biệt", "Bỏ luôn", "Bỏ khi chỉ có Cover; Giữ khi có Audio"],
        ["Spotify / Ghép chéo", "Đề xuất crawl", "Thấy không ổn", "Phản biện: Lệch domain + Leakage"],
        ["Cơ chế Confidence", "Chưa có", "Thêm confidence", "4 vai trò: Fusion, Loss, Pseudo, Review"],
        ["SSL Pretraining", "Chưa rõ", "Train nhận diện trước", "Gọi tên chuẩn: AudioMAE / CLAP"],
        ["Ngưỡng Pseudo-label", "Chưa đề cập", "Ngưỡng cố định 95%", "Adaptive Threshold (0.68 - 0.88) + Top-K"],
        ["Phân loại Thể loại", "Multi-class (Softmax)", "Multi-label (Sigmoid)", "Multi-label + Asymmetric Loss + Tau_c"]
    ]
    create_table(col_w_comp, headers_comp, data_comp, header_bg="1B365D")

    # ==========================================
    # PHẦN 5: CẨM NANG TÓM TẮT TOÀN DIỆN
    # ==========================================
    add_heading_1("PHẦN 5: CẨM NANG TÓM TẮT TOÀN BỘ GIẢI PHÁP (BẢO VỆ TRƯỚC HỘI ĐỒNG)")

    add_heading_2("5.1. Bảng Tổng Hợp Xử Lý Modality Thiếu")
    headers_ms = ["Tình huống dữ liệu", "Giải pháp xử lý chuẩn mực", "Lưu ý then chốt"]
    col_w_ms = [Inches(1.8), Inches(2.7), Inches(2.0)]
    data_ms = [
        ["Thiếu lyric, có audio + cover", "Tách vocal bằng Demucs -> ASR (PhoWhisper) -> cờ is_asr=True; Hoặc gán mask_lyric=0.", "Gán confidence = 0.6 cho lyric ASR."],
        ["Thiếu audio", "mask_audio = 0, chỉ dùng lyric + cover.", "Nếu bài toán coi audio là bắt buộc -> Loại bỏ."],
        ["Thiếu cover", "mask_cover = 0, tăng modality dropout (0.4 - 0.5).", "Cover nhạc Việt chứa nhiều ảnh ca sĩ gây overfit."],
        ["Chỉ có cover", "Bỏ hoàn toàn khỏi tập train.", "Không đủ tín hiệu âm học."],
        ["Nhạc cổ truyền", "KHÔNG DÙNG ASR. Gán mask_lyric = 0.", "Nhận diện 100% bằng âm thanh (Audio-dominant)."]
    ]
    create_table(col_w_ms, headers_ms, data_ms, header_bg="005A9E")

    add_heading_2("5.2. Quy Trình 3 Cấp Độ Xử Lý Genre Thiếu (368k bài)")
    add_bullet("Fuzzy matching theo (title, artist) với Levenshtein >= 0.92 để lấy lại ~3.200 nhãn từ nguồn có nhãn.", "Cấp 1 — Cross-Dedup Recovery: ");
    add_bullet("Huấn luyện AudioMAE / CLAP trên 355k bài không nhãn để mô hình hiểu đặc trưng âm sắc nhạc Việt.", "Cấp 2 — SSL Pretraining: ");
    add_bullet("Dùng Teacher gán nhãn với Adaptive Threshold (Pop >= 0.88, Nhạc Đỏ >= 0.68), lọc Top 50k bài huấn luyện Student kèm lambda = 0.5 và Label Smoothing = 0.1.", "Cấp 3 — Controlled Pseudo-Labeling: ");

    add_heading_2("5.3. Xử Lý Nhiều Genre (Multi-Label Classification)")
    add_bullet("Thay Softmax bằng 8 hàm Sigmoid độc lập. Vector Multi-hot y thuộc {0, 1}^8 (Ví dụ: Pop + R&B -> [1, 0, 0, 0, 0, 0, 1, 0]).", "Formulation: ");
    add_bullet("Dùng Asymmetric Loss (ASL) để dập tắt gradient từ các nhãn 0 dễ đoán, giải quyết tỷ lệ mất cân bằng âm/dương 7:1.", "Loss Function: ");
    add_bullet("Quét tìm ngưỡng quyết định tối ưu riêng tau_c thuộc [0.2, 0.8] trên tập Validation thay vì dùng cố định 0.5.", "Decision Threshold: ");

    add_heading_2("5.4. Chiến Lược Train Trên GPU 4GB VRAM (RTX 3050 Ti Laptop)")
    add_bullet("Trích xuất offline vector biểu diễn: Audio -> PANNs CNN14 (512-D), Lyric -> PhoBERT (768-D), Cover -> CLIP ViT (512-D), lưu thành file Parquet / NPY.", "Offline Feature Store: ");
    add_bullet("Khi train chỉ nạp các vector float32 vào mạng Fusion MLP siêu nhẹ. VRAM tiêu thụ chưa đầy 350MB, tốc độ: 100 epoch chỉ mất 45 giây!", "Ultra-fast Fusion: ");
    add_bullet("Áp dụng Modality Dropout ngẫu nhiên (tắt nhánh audio, lyric hoặc cover) để ép mạng không bị phụ thuộc độc quyền vào Audio.", "Chống Audio Dominance: ");

    add_heading_2("5.5. Đánh Giá & Chia Dữ Liệu Chuẩn NCKH")
    add_bullet("Macro F1-score (công bằng lớp hiếm), Micro F1-score, mAP, Subset Accuracy (Exact Match).", "Metrics: ");
    add_bullet("Bắt buộc chia tập bằng GroupShuffleSplit theo artist_id để tránh hiện tượng rò rỉ ca sĩ (Artist Leakage).", "Data Split: ");

    # ==========================================
    # CÂU CHỐT BẢO VỆ
    # ==========================================
    add_heading_1("CÂU CHỐT BẢO VỆ TOÀN DIỆN TRƯỚC HỘI ĐỒNG NCKH")
    add_callout("CÂU TRẢ LỜI ĐẠT ĐIỂM XUẤT SẮC TRƯỚC HỘI ĐỒNG",
        "\"Thưa Hội đồng, để giải quyết bài toán nhạc Việt trên dữ liệu khuyết thiếu thực tế:\n"
        "1. Với Modality khuyết: Chúng em không bịa dữ liệu giả; chúng em tích hợp cơ chế Mask-Aware Gated Fusion kết hợp Confidence per feature. Trong đó, tách vocal trước khi ASR cho nhạc trẻ, loại trừ hoàn toàn ASR trên nhạc cổ truyền, và không dùng ảnh ngoài để tránh rò rỉ ca sĩ.\n"
        "2. Với Genre khuyết: Chúng em áp dụng quy trình 3 cấp độ: Khôi phục bằng khử trùng lặp -> Học tự giám sát AudioMAE trên 355k bài Zalo AI -> Gán nhãn giả bán giám sát có kiểm soát bằng Adaptive Threshold theo từng lớp.\n"
        "3. Với bài nhiều Genre: Chúng em chuẩn hóa về bài toán Multi-label với 8 đầu ra Sigmoid độc lập và tối ưu bằng hàm Asymmetric Loss (ASL) để giải quyết triệt để mất cân bằng mẫu âm/dương.\n"
        "4. Với giới hạn phần cứng 4GB VRAM: Chúng em triển khai kiến trúc Two-Stage Offline Feature Store, nạp vector nhúng thay vì nạp backbone, giúp hệ thống chạy mượt mà dưới 400MB VRAM mà không bị OOM.\"")

    # Save to dedicated clean file
    output_path = "f:/Data/reports/NHAN_XET_VA_TONG_KET_HOAN_THIEN_MULTIMODAL_GENRE.docx"
    doc.save(output_path)
    print(f"Report generated successfully at: {output_path}")

if __name__ == "__main__":
    create_review_docx()
